"""
Report generators: Excel (openpyxl) and PDF (reportlab).
"""
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

HEADER_FILL  = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')
HEADER_FONT  = Font(color='FFFFFF', bold=True)
ALT_FILL     = PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid')
THIN         = Side(border_style='thin', color='CCCCCC')
CELL_BORDER  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header_row(ws, row, col_count):
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = CELL_BORDER


def _style_data_row(ws, row, col_count, alt=False):
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        if alt:
            cell.fill = ALT_FILL
        cell.alignment = Alignment(vertical='center')
        cell.border = CELL_BORDER


def _auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)


# ─── Daily Attendance Report ──────────────────────────────────────────────────

def generate_daily_excel(organization, target_date: date) -> bytes:
    from apps.attendance.models import AttendanceRecord

    wb = Workbook()
    ws = wb.active
    ws.title = f"Daily_{target_date}"

    headers = ['Employee ID', 'Name', 'Department', 'Check In', 'Check Out',
               'Working Hours', 'Status', 'Late (min)', 'Method']
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    records = (
        AttendanceRecord.objects
        .filter(organization=organization, date=target_date)
        .select_related('employee__user', 'employee__department')
        .order_by('employee__user__last_name')
    )

    for i, r in enumerate(records, start=2):
        emp = r.employee
        ws.append([
            emp.employee_id,
            emp.user.get_full_name(),
            emp.department.name if emp.department else '',
            r.check_in_time.strftime('%H:%M') if r.check_in_time else '',
            r.check_out_time.strftime('%H:%M') if r.check_out_time else '',
            str(r.working_hours or ''),
            r.get_status_display(),
            r.late_minutes or 0,
            r.get_check_in_method_display() if r.check_in_method else '',
        ])
        _style_data_row(ws, i, len(headers), alt=(i % 2 == 0))

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_monthly_excel(organization, year: int, month: int) -> bytes:
    from apps.attendance.models import AttendanceRecord
    from apps.accounts.models import CustomUser
    import calendar

    _, days_in_month = calendar.monthrange(year, month)
    day_range = range(1, days_in_month + 1)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}"

    # Build header: Employee ID | Name | Dept | Day1 … DayN | P | A | L | HD | %
    day_headers = [str(d) for d in day_range]
    headers = ['ID', 'Name', 'Department'] + day_headers + ['P', 'A', 'L', 'HD', '%']
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    employees = (
        CustomUser.objects
        .filter(organization=organization, is_active=True)
        .exclude(role='student')
        .select_related('department')
        .order_by('last_name')
    )

    STATUS_MAP = {'present': 'P', 'absent': 'A', 'on_leave': 'L',
                  'half_day': 'HD', 'late': 'P*', 'holiday': 'H',
                  'weekend': '-', 'regularized': 'R'}

    for row_idx, user in enumerate(employees, start=2):
        try:
            emp = user.employee_profile
        except Exception:
            continue

        records = {
            r.date.day: r
            for r in AttendanceRecord.objects.filter(
                employee=emp, date__year=year, date__month=month
            )
        }

        present = absent = leave = half_day = 0
        day_cells = []
        for d in day_range:
            r = records.get(d)
            if r:
                code = STATUS_MAP.get(r.status, r.status)
                if r.status == 'present': present += 1
                elif r.status == 'absent': absent += 1
                elif r.status == 'on_leave': leave += 1
                elif r.status == 'half_day': half_day += 1
                elif r.status == 'late': present += 1
            else:
                code = ''
            day_cells.append(code)

        total = present + absent + leave + half_day
        pct = round(present / total * 100, 1) if total else 0

        row_data = [
            emp.employee_id,
            user.get_full_name(),
            user.department.name if user.department else '',
        ] + day_cells + [present, absent, leave, half_day, f"{pct}%"]
        ws.append(row_data)
        _style_data_row(ws, row_idx, len(headers), alt=(row_idx % 2 == 0))

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_department_excel(organization, start_date: date, end_date: date) -> bytes:
    from apps.attendance.models import AttendanceRecord
    from apps.organizations.models import Department

    wb = Workbook()

    departments = Department.objects.filter(organization=organization)
    for dept in departments:
        ws = wb.create_sheet(title=dept.name[:31])
        headers = ['Employee ID', 'Name', 'Total Days', 'Present', 'Absent',
                   'Leave', 'Half Day', 'Late', 'Attendance %']
        ws.append(headers)
        _style_header_row(ws, 1, len(headers))

        employees = dept.members.filter(is_active=True).select_related('employee_profile')
        for row_idx, user in enumerate(employees, start=2):
            try:
                emp = user.employee_profile
            except Exception:
                continue
            records = AttendanceRecord.objects.filter(
                employee=emp, date__range=[start_date, end_date]
            )
            total   = records.count()
            present = records.filter(status__in=['present', 'late']).count()
            absent  = records.filter(status='absent').count()
            leave   = records.filter(status='on_leave').count()
            half_d  = records.filter(status='half_day').count()
            late    = records.filter(status='late').count()
            pct     = round(present / total * 100, 1) if total else 0
            ws.append([
                emp.employee_id, user.get_full_name(),
                total, present, absent, leave, half_d, late, f"{pct}%"
            ])
            _style_data_row(ws, row_idx, len(headers), alt=(row_idx % 2 == 0))
        _auto_width(ws)

    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_payroll_excel(organization, year: int, month: int) -> bytes:
    from apps.attendance.models import AttendanceRecord
    from apps.accounts.models import CustomUser

    wb = Workbook()
    ws = wb.active
    ws.title = f"Payroll_{year}-{month:02d}"

    headers = ['Employee ID', 'Name', 'Department', 'Working Days',
               'Present Days', 'Absent Days', 'Leave Days', 'Half Days',
               'Total Hours', 'Overtime Hours', 'Late Minutes']
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    employees = (
        CustomUser.objects
        .filter(organization=organization, is_active=True)
        .exclude(role='student')
        .select_related('department')
        .order_by('last_name')
    )

    for row_idx, user in enumerate(employees, start=2):
        try:
            emp = user.employee_profile
        except Exception:
            continue
        records = AttendanceRecord.objects.filter(
            employee=emp, date__year=year, date__month=month
        )
        total_hours    = sum((r.working_hours or 0) for r in records)
        overtime_hours = sum((r.overtime_hours or 0) for r in records)
        late_mins      = sum((r.late_minutes or 0) for r in records)
        ws.append([
            emp.employee_id,
            user.get_full_name(),
            user.department.name if user.department else '',
            records.exclude(status__in=['holiday', 'weekend']).count(),
            records.filter(status__in=['present', 'late']).count(),
            records.filter(status='absent').count(),
            records.filter(status='on_leave').count(),
            records.filter(status='half_day').count(),
            float(total_hours),
            float(overtime_hours),
            int(late_mins),
        ])
        _style_data_row(ws, row_idx, len(headers), alt=(row_idx % 2 == 0))

    _auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── PDF generators ───────────────────────────────────────────────────────────

def generate_daily_pdf(organization, target_date: date) -> bytes:
    from apps.attendance.models import AttendanceRecord

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=10*mm, rightMargin=10*mm,
                            topMargin=10*mm, bottomMargin=10*mm)

    styles  = getSampleStyleSheet()
    title_s = ParagraphStyle('Title', fontSize=14, fontName='Helvetica-Bold',
                             alignment=TA_CENTER, spaceAfter=6)
    sub_s   = ParagraphStyle('Sub', fontSize=10, alignment=TA_CENTER, spaceAfter=12)

    elements = [
        Paragraph(f"{organization.name} — Daily Attendance", title_s),
        Paragraph(str(target_date), sub_s),
    ]

    records = (
        AttendanceRecord.objects
        .filter(organization=organization, date=target_date)
        .select_related('employee__user', 'employee__department')
        .order_by('employee__user__last_name')
    )

    data = [['Emp ID', 'Name', 'Department', 'In', 'Out', 'Hours', 'Status', 'Late(m)']]
    for r in records:
        emp = r.employee
        data.append([
            emp.employee_id,
            emp.user.get_full_name(),
            emp.department.name if emp.department else '',
            r.check_in_time.strftime('%H:%M') if r.check_in_time else '-',
            r.check_out_time.strftime('%H:%M') if r.check_out_time else '-',
            str(r.working_hours or '-'),
            r.get_status_display(),
            str(r.late_minutes or 0),
        ])

    col_widths = [22*mm, 45*mm, 35*mm, 16*mm, 16*mm, 18*mm, 22*mm, 16*mm]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND',  (0, 0), (-1, 0), colors.HexColor('#1565C0')),
        ('TEXTCOLOR',   (0, 0), (-1, 0), colors.white),
        ('FONTNAME',    (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE',    (0, 0), (-1, -1), 8),
        ('ALIGN',       (0, 0), (-1, -1), 'CENTER'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#E3F2FD')]),
        ('GRID',        (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('VALIGN',      (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(t)
    doc.build(elements)
    return buf.getvalue()
